from docx import Document
from rfp_copilot.ingestion.loaders import load_docx, load_document
import pytest


def test_docx_loader_preserves_heading_paragraph_and_table_order(tmp_path):
    path = tmp_path / "rfp.docx"
    document = Document()
    document.add_heading("Scope", level=1)
    document.add_paragraph("The vendor shall provide SSO.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Requirement"
    table.rows[0].cells[1].text = "Priority"
    table.rows[1].cells[0].text = "Encryption"
    table.rows[1].cells[1].text = "Mandatory"
    document.save(path)

    loaded = load_docx(str(path))
    assert loaded.blocks[0].text == "## Scope"
    assert "SSO" in loaded.blocks[1].text
    assert loaded.blocks[2].block_type == "table"
    assert "Requirement: Encryption" in loaded.blocks[2].text


def test_loader_rejects_missing_and_unsupported_files(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_document(str(tmp_path / "missing.pdf"))
    unsupported = tmp_path / "rfp.txt"
    unsupported.write_text("text")
    with pytest.raises(ValueError):
        load_document(str(unsupported))
