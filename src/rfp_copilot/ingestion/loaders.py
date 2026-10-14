"""Page- and block-aware PDF/DOCX loading for the Week 5 ingestion pipeline."""

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


@dataclass
class DocumentBlock:
    """One source-ordered piece of a document, retaining its provenance."""

    text: str
    page: int | None = None
    section: str | None = None
    block_index: int = 0
    block_type: str = "paragraph"


@dataclass
class LoadedDocument:
    text: str
    pages: list[str]
    blocks: list[DocumentBlock] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def _table_to_text(table: Table) -> str:
    """Turn a DOCX table into readable header/value rows without losing values."""
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    headers = rows[0]
    result = []
    for values in rows[1:]:
        pairs = [f"{header}: {value}" for header, value in zip(headers, values) if value]
        if pairs:
            result.append("; ".join(pairs))
    return "\n".join(result) if result else " | ".join(headers)


def load_pdf(path: str) -> LoadedDocument:
    """Extract each PDF page separately so page provenance remains available."""
    reader = PdfReader(path)
    pages = []
    blocks = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(text)
        if text.strip():
            blocks.append(DocumentBlock(text=text, page=index, block_index=index - 1, block_type="page"))
    return LoadedDocument(
        text="\n\n".join(pages),
        pages=pages,
        blocks=blocks,
        metadata={"source_type": "pdf", "filename": Path(path).name, "page_count": len(pages)},
    )


def load_docx(path: str) -> LoadedDocument:
    """Read paragraphs and tables in native body order, preserving headings."""
    document = Document(path)
    blocks = []
    current_section = None
    table_count = 0
    paragraph_count = 0

    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.startswith("Heading"):
                current_section = text
                text = f"## {text}"
            blocks.append(DocumentBlock(text=text, section=current_section, block_index=len(blocks)))
            paragraph_count += 1
        elif element.tag.endswith("}tbl"):
            text = _table_to_text(Table(element, document))
            if text:
                blocks.append(DocumentBlock(text=f"[TABLE]\n{text}", section=current_section, block_index=len(blocks), block_type="table"))
            table_count += 1

    block_text = [block.text for block in blocks]
    return LoadedDocument(
        text="\n\n".join(block_text),
        pages=block_text,
        blocks=blocks,
        metadata={
            "source_type": "docx",
            "filename": Path(path).name,
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "heading_count": sum(1 for block in blocks if block.text.startswith("## ")),
        },
    )


def load_document(path: str) -> LoadedDocument:
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"Target document not found: {path}")
    if file_path.suffix.lower() == ".pdf":
        return load_pdf(str(file_path))
    if file_path.suffix.lower() == ".docx":
        return load_docx(str(file_path))
    raise ValueError(f"Unsupported file type: {path}")
